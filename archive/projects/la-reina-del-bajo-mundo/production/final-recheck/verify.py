from pathlib import Path
from copy import copy
import json,re,hashlib
import openpyxl
from docx import Document
base=Path(__file__).resolve().parent
root=base.parent
old=openpyxl.load_workbook('/Users/apple/Desktop/前夫黑道酵母_分镜_21-33.xlsx')
new=openpyxl.load_workbook(base/'checked-workbook.xlsx')
assert old.sheetnames==new.sheetnames==['分镜脚本']
a,b=old.active,new.active
norm=lambda v:None if v=='' else v
changes=[]
rows=[]
for folder in ['production-21-25','production-26-30','production-31-33']:
 rows.extend(json.loads((root/folder/'rows.json').read_text()))
for ri in range(1,61):
 for ci in range(1,10):
  x,y=a.cell(ri,ci),b.cell(ri,ci)
  if norm(x.value)!=norm(y.value):
   assert ri>1 and ci in [6,9],(ri,ci,'unintended original or asset edit')
   changes.append({'cell':y.coordinate,'storyboard':b.cell(ri,2).value})
  for prop in ['font','alignment','border','number_format']:
   assert copy(getattr(x,prop))==copy(getattr(y,prop)),(ri,ci,prop)
  assert x.fill.patternType==y.fill.patternType
  assert x.fill.fgColor.rgb==y.fill.fgColor.rgb
  if ri>1:assert norm(y.value)==norm(rows[ri-2][ci-1])
 assert a.row_dimensions[ri].height==b.row_dimensions[ri].height
for c in 'ABCDEFGHI':assert a.column_dimensions[c].width==b.column_dimensions[c].width
assert a.freeze_panes==b.freeze_panes=='A2'
assert not b.merged_cells.ranges
source=[(i+1,p.text) for i,p in enumerate(Document(root/'repository/archive/projects/la-reina-del-bajo-mundo/sources/2026-10-29/第21集-33集.docx').paragraphs) if p.text.strip()]
corrected=[text.replace('Vera（','Renata（',1) if i==135 else text for i,text in source]
assert '\n'.join(r[7] for r in rows)=='\n'.join(corrected)
rx=re.compile(r'^([A-Za-zÁÉÍÓÚÜÑáéíóúüñ][A-Za-zÁÉÍÓÚÜÑáéíóúüñ ]*?)(?:\s+(VO|OS))?(?:（[^）]*）)?：(.*)$',re.S)
expected=[]
for (i,_),text in zip(source,corrected):
 m=rx.match(text.strip())
 if m:expected.append((i,m.group(1),m.group(2) or '对白',m.group(3)))
mapping=[]
for name in ['21-25','26-30','31-33']:mapping+=json.loads((root/f'production-{name}'/'speech_mapping.json').read_text())
assert expected==[(v['source_index'],v['name'],v['kind'],v['text']) for v in mapping]
for v in mapping:
 row=next(r for r in rows if r[1]==v['storyboard'])
 assert row[5].count('「'+v['text']+'」')==1
 assert v['end']<=float(row[8][:-1])
assert all(float(t)<=4 for row in rows for t in re.findall(r'【时长】([0-9.]+)s',row[5]))
assert all(float(row[8][:-1])<=30 for row in rows)
assert not any(c.data_type=='f' for row in b for c in row)
report={'result':'passed','episodes':list(range(21,34)),'rows':len(rows),'shots':sum(int(re.search(r'\[运镜时间线（(\d+)个硬切分镜）\]',r[5]).group(1)) for r in rows),'seconds':sum(float(r[8][:-1]) for r in rows),'dialogue':108,'VO':10,'OS':1,'raw_source_blocks':len(source),'source_text_and_punctuation_preserved':True,'speech_source_kind_order_count_and_punctuation_preserved':True,'asset_columns_preserved':True,'styles_widths_heights_freeze_and_sheet_count_preserved':True,'changed_cells':changes,'video_validation':'未实测','sha256':hashlib.sha256((base/'checked-workbook.xlsx').read_bytes()).hexdigest()}
(base/'RECHECK_RESULT.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in report.items() if k!='changed_cells'},ensure_ascii=False))
