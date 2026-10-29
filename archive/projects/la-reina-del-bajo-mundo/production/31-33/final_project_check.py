from pathlib import Path
import json,re,hashlib
from docx import Document
import openpyxl
base=Path(__file__).resolve().parent
project=base.parent/'repository/archive/projects/la-reina-del-bajo-mundo'
source=json.loads((project/'SCRIPT_SOURCE.json').read_text())['blocks']
raw=Document(project/'sources/2026-10-29/第21集-33集.docx')
raw_blocks=[{'body_index':i+1,'text':p.text} for i,p in enumerate(raw.paragraphs) if p.text.strip()]
assert [(x['body_index'],x['text']) for x in raw_blocks]==[(x['body_index'],x['text']) for x in source],'raw DOCX differs from extracted baseline'
wb=openpyxl.load_workbook(base/'checked-workbook.xlsx')
sheet=wb.active
rows=[list(r) for r in sheet.iter_rows(min_row=2,max_col=9,values_only=True) if r[1] is not None]
assert len(rows)==59
assert set(row[0] for row in rows)==set(range(21,34))
assert len(set(row[1] for row in rows))==59
corrected=[]
for block in raw_blocks:
 text=block['text']
 if block['body_index']==135:text=text.replace('Vera（','Renata（',1)
 corrected.append(text)
assert '\n'.join(row[7] for row in rows)=='\n'.join(corrected),'original text coverage/order differs'
rx=re.compile(r'^([A-Za-zÁÉÍÓÚÜÑáéíóúüñ][A-Za-zÁÉÍÓÚÜÑáéíóúüñ ]*?)(?:\s+(VO|OS))?(?:（[^）]*）)?：(.*)$',re.S)
expected=[]
for block,text in zip(raw_blocks,corrected):
 m=rx.match(text.strip())
 if m:expected.append((block['body_index'],m.group(1),m.group(2) or '对白',m.group(3)))
actual=[]
for folder in [base.parent/'production-21-25',base.parent/'production-26-30',base]:
 actual.extend(json.loads((folder/'speech_mapping.json').read_text()))
assert [(x['source_index'],x['name'],x['kind'],x['text']) for x in actual]==expected,'speech baseline mismatch'
byid={row[1]:row for row in rows}
for item in actual:
 prompt=byid[item['storyboard']][5]
 assert prompt.count('「'+item['text']+'」')==1,(item['source_index'],'prompt voice occurrence')
 assert item['end']<=float(byid[item['storyboard']][8][:-1])
assert all(not c.data_type=='f' for row in sheet for c in row),'unexpected formulas'
dur=sum(float(row[8][:-1]) for row in rows)
shots=sum(int(re.search(r'\[运镜时间线（(\d+)个硬切分镜）\]',row[5]).group(1)) for row in rows)
assert shots==319 and dur==1032.5
assert all(float(t)<=4 for row in rows for t in re.findall(r'【时长】([0-9.]+)s',row[5]))
assert all(float(row[8][:-1])<=30 for row in rows)
report={'result':'passed','episodes':list(range(21,34)),'source_nonempty_blocks':len(raw_blocks),'raw_docx_equals_baseline':True,'original_column_complete_in_order':True,'approved_source_corrections':[135],'storyboards':len(rows),'shots':shots,'seconds':dur,'dialogue':sum(x[2]=='对白' for x in expected),'VO':sum(x[2]=='VO' for x in expected),'OS':sum(x[2]=='OS' for x in expected),'speech_occurrences':len(expected),'exact_speech_source_kind_order_and_punctuation':True,'each_prompt_voice_occurrence_once':True,'video_validation':'未实测','sha256':hashlib.sha256((base/'checked-workbook.xlsx').read_bytes()).hexdigest()}
(base/'PROJECT_FINAL_CHECK.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
