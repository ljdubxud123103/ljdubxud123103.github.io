from pathlib import Path
from copy import copy
import openpyxl,json,re,hashlib
b=Path(__file__).resolve().parent
target=b/'before-workbook.xlsx'
old=openpyxl.load_workbook(target);new=openpyxl.load_workbook(b/'checked-workbook.xlsx')
assert old.sheetnames==new.sheetnames==['分镜脚本']
a,c=old.active,new.active
before=[list(r) for r in a.iter_rows(min_row=2,max_col=9,values_only=True) if r[1] is not None]
after=[list(r) for r in c.iter_rows(min_row=2,max_col=9,values_only=True) if r[1] is not None]
expected=json.loads((b/'rows.json').read_text())
norm=lambda x:None if x=='' else x
assert [[norm(x) for x in r] for r in after]==[[norm(x) for x in r] for r in expected]
assert len(after)==len(before)-1==58
assert '\n'.join(r[7] for r in before)=='\n'.join(r[7] for r in after)
unchanged=0
for r in before:
 if r[1] not in ['81-2A','81-3A','81-4A']:
  other=next(x for x in after if x[1]==r[1]);assert list(map(norm,r))==list(map(norm,other));unchanged+=1
assert not any(r[1] in ['81-3A','81-4A'] for r in after)
for ri,r in enumerate(after,2):
 assert c.row_dimensions[ri].height==a.row_dimensions[2].height
 for ci in range(1,10):
  for p in ['font','border','alignment','number_format']:
   assert copy(getattr(c.cell(ri,ci),p))==copy(getattr(a.cell(2,ci),p))
assert c.freeze_panes==a.freeze_panes
assert [x.value for x in c[1]]==[x.value for x in a[1]]
for col in 'ABCDEFGHI':assert a.column_dimensions[col].width==c.column_dimensions[col].width
for id,duration in [('81-2A',30),('81-2B',18)]:
 r=next(r for r in after if r[1]==id)
 times=list(map(float,re.findall(r'【时长】([0-9.]+)s',r[5])))
 assert sum(times)==duration and r[8]==str(duration)+'s'
 assert r[6] is None
 assert max(times)<=4 or id=='81-2A' and times==[3,3,4,4,4,8,4]
 if id=='81-2A':assert '18s—26s，真正连续一镜到底' in r[5] and '0—1秒' in r[5] and '1—6.5秒' in r[5] and '6.5—8秒' in r[5]
 if id=='81-2B':assert '一个小红点' in r[5] and '空气中不出现射线' in r[5] and '细红光' not in r[5]
state=json.loads((b/'PLAN_CHECK.json').read_text())['end_start_state']
assert state in next(r for r in after if r[1]=='81-2A')[5]
assert state in next(r for r in after if r[1]=='81-2B')[5].split('[运镜时间线]')[0]
report={'result':'passed','old_storyboards':['81-2A','81-3A','81-4A'],'new_storyboards':['81-2A','81-2B'],'seconds':[30,18],'walking':'第一段18—26秒连续全身跟拍，一镜到底，无内部硬切','source_column_preserved':True,'other_storyboards_unchanged':unchanged,'format_and_freeze_preserved':True,'new_segment_state_identical':True,'laser_policy':'人物表面单一小红点，无空气射线','video_validation':'未实测','sha256':hashlib.sha256((b/'checked-workbook.xlsx').read_bytes()).hexdigest()}
(b/'WORKBOOK_CHECK.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
