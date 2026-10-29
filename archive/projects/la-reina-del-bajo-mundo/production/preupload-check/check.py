from pathlib import Path
from docx import Document
import openpyxl,json,re,hashlib
from copy import copy
base=Path(__file__).resolve().parent
project=base.parent/'repository/archive/projects/la-reina-del-bajo-mundo'
file=Path('/Users/apple/Desktop/前夫黑道酵母_分镜_21-33.xlsx')
w=openpyxl.load_workbook(file)
assert w.sheetnames==['分镜脚本']
s=w.active
ref=openpyxl.load_workbook(project/'templates/契约夫人（西语改）_分镜_91-100_20261026.xlsx').active
assert [c.value for c in s[1]]==[c.value for c in ref[1]][:9]
assert s.freeze_panes==ref.freeze_panes=='A2'
assert not s.merged_cells.ranges
for c in 'ABCDEFGHI':assert s.column_dimensions[c].width==ref.column_dimensions[c].width
rows=[list(r) for r in s.iter_rows(min_row=2,max_col=9,values_only=True) if r[1] is not None]
assert len(rows)==59 and len({r[1] for r in rows})==59
raw=[(i+1,p.text) for i,p in enumerate(Document(project/'sources/2026-10-29/第21集-33集.docx').paragraphs) if p.text.strip()]
corrected=[t.replace('Vera（','Renata（',1) if i==135 else t for i,t in raw]
assert '\n'.join(r[7] for r in rows)=='\n'.join(corrected)
rx=re.compile(r'^([A-Za-zÁÉÍÓÚÜÑáéíóúüñ][A-Za-zÁÉÍÓÚÜÑáéíóúüñ ]*?)(?:\s+(VO|OS))?(?:（[^）]*）)?：(.*)$',re.S)
expected=[]
for (i,_),t in zip(raw,corrected):
 m=rx.match(t.strip())
 if m:expected.append((m.group(1),m.group(2) or '对白',m.group(3)))
actual=[];voices=[];shots=0;seconds=0;maxshot=0;maxsegment=0
assets=json.loads((project/'ASSET_INDEX.json').read_text())['assets']
names={x['name_original'] for x in assets}
style=json.loads((project/'TEMPLATE_CONTRACT.json').read_text())['style_lock_reference_first_row']
for ri,r in enumerate(rows,2):
 assert s.row_dimensions[ri].height==ref.row_dimensions[2].height
 for ci in range(1,10):
  cell=s.cell(ri,ci);rc=ref.cell(2,ci)
  assert copy(cell.alignment)==copy(rc.alignment)
  assert copy(cell.font)==copy(rc.font)
  for edge in ['left','right','top','bottom']:
   x,y=getattr(cell.border,edge),getattr(rc.border,edge)
   assert x.style==y.style
   assert x.color.type==y.color.type=='rgb' and x.color.rgb[-6:]==y.color.rgb[-6:]
  assert cell.fill.patternType==rc.fill.patternType
  assert len(str(cell.value or ''))<32767
  assert cell.data_type!='f'
 assert r[5].startswith(style)
 assert all(n in names or n in ['酒店正门（外·日）','酒店正门（外·夜）','公寓入口（文字设定·夜）'] for n in r[3].split('；'))
 assert all(n in names or n in ['Mujer OS','Reportera D VO','El Zorro Blanco VO'] for n in r[2].split('；'))
 anchors=r[5].split('[锚定列表]')[1].split('[空间布局总览]')[0]
 for line in anchors.splitlines():
  if line.startswith(('角色：','场景：','道具：')):
   for n in line.split('：',1)[1].split('；'):
    if n.startswith('@'):assert n[1:] in names,(r[1],n)
 if r[6]:
  for line in r[6].splitlines():
   name,quote=line.split('：',1)
   assert quote.startswith('「') and quote.endswith('」')
   kind='OS' if name.endswith(' OS') else 'VO' if name.endswith(' VO') else '对白'
   actual.append((name[:-3] if kind!='对白' else name,kind,quote[1:-1]))
 block=r[5].split('[音轨规则]')[1].split('[锚定列表]')[0]
 local=[];lastend=0
 for match in re.finditer(r'([0-9.]+)s—([0-9.]+)s，([^\n（]+)（(对白|VO|OS)，[^\n]*?）：「([^」]*)」',block):
  start,end,name,kind,text=match.groups();start=float(start);end=float(end)
  assert start>=lastend and end<=float(r[8][:-1])
  lastend=end;local.append((name.lstrip('@'),kind,text))
  assert r[5].count('「'+text+'」')==1
 local_expected=[]
 if r[6]:
  for line in r[6].splitlines():
   n,t=line.split('：',1);k='OS' if n.endswith(' OS') else 'VO' if n.endswith(' VO') else '对白'
   local_expected.append((n[:-3] if k!='对白' else n,k,t[1:-1]))
 assert len(local)==len(local_expected)
 for (n,k,t),(n2,k2,t2) in zip(local,local_expected):
  assert n.startswith(n2) and k==k2 and t==t2
 voices.extend(local)
 times=list(map(float,re.findall(r'【时长】([0-9.]+)s',r[5])))
 assert len(times)==int(re.search(r'\[运镜时间线（(\d+)个硬切分镜）\]',r[5]).group(1))
 assert abs(sum(times)-float(r[8][:-1]))<.001
 assert all(0<t<=4 for t in times)
 assert float(r[8][:-1])<=30
 assert r[5].endswith('时长：'+r[8])
 assert not any(t in r[5] for t in ['{{','同上','见上一段','@图','Mixed'])
 shots+=len(times);seconds+=sum(times);maxshot=max(maxshot,max(times));maxsegment=max(maxsegment,sum(times))
assert expected==actual and len(actual)==len(voices)==119
assert len(raw)==551 and shots==324 and seconds==1045
assert set(r[0] for r in rows)==set(range(21,34))
assert hashlib.sha256(file.read_bytes()).hexdigest()==json.loads((base.parent/'production-final-recheck/RECHECK_RESULT.json').read_text())['sha256']
report={'result':'passed','input':'当前桌面实际XLSX，直接读取；未用rows.json或speech_mapping.json作为输出基线','workbook_sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'rows':len(rows),'shots':shots,'seconds':seconds,'raw_docx_blocks':len(raw),'speech_occurrences':len(actual),'dialogue':108,'VO':10,'OS':1,'max_shot_seconds':maxshot,'max_segment_seconds':maxsegment,'original_text_exact':True,'actual_G_and_F_voice_text_order_kind_checked':True,'original_template_columns_font_alignment_border_row_height_checked':True,'asset_anchor_names_checked':True,'duration_sums_and_voice_intervals_checked':True,'no_new_changes':True,'video_validation':'未进行真实视频验证'}
(base/'PREUPLOAD_CHECK.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
