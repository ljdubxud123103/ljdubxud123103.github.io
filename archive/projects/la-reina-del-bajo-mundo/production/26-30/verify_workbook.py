from pathlib import Path
from copy import copy
import json,hashlib,re
import openpyxl
BASE=Path(__file__).resolve().parent
project=BASE.parent/'repository/archive/projects/la-reina-del-bajo-mundo'
target=Path('/Users/apple/Desktop/前夫黑道酵母_分镜_21-33.xlsx')
old=openpyxl.load_workbook(target)
new=openpyxl.load_workbook(BASE/'checked-workbook.xlsx')
assert old.sheetnames==new.sheetnames==['分镜脚本']
a,b=old.active,new.active
norm=lambda v:None if v=='' else v
for row in range(1,26):
 for col in range(1,10):
  assert norm(a.cell(row,col).value)==norm(b.cell(row,col).value),(row,col,'previous content')
  for prop in ['font','alignment','border','number_format']:
   assert copy(getattr(a.cell(row,col),prop))==copy(getattr(b.cell(row,col),prop)),(row,col,prop)
assert a.freeze_panes==b.freeze_panes=='A2'
assert not b.merged_cells.ranges
for c in 'ABCDEFGHI':assert a.column_dimensions[c].width==b.column_dimensions[c].width
rows=json.loads((BASE/'rows.json').read_text())
for ri,row in enumerate(rows,26):
 for ci,value in enumerate(row,1):
  assert norm(b.cell(ri,ci).value)==norm(value),(ri,ci,'new content')
  for prop in ['font','alignment','border','number_format']:
   assert copy(getattr(a.cell(2,ci),prop))==copy(getattr(b.cell(ri,ci),prop)),(ri,ci,prop)
 assert b.row_dimensions[ri].height==a.row_dimensions[2].height
assert max(c.row for row in b for c in row if c.value is not None)==46
assert max(c.column for row in b for c in row if c.value is not None)==9
known={x['name_original'] for x in json.loads((project/'ASSET_INDEX.json').read_text())['assets']}
for row in rows:
 assert all(n in known for n in row[3].split('；')),(row[1],'scene asset')
 assert all(n in known or n in ['Reportera D VO','El Zorro Blanco VO'] for n in row[2].split('；'))
 assert row[5].startswith(json.loads((project/'TEMPLATE_CONTRACT.json').read_text())['style_lock_reference_first_row'])
 assert max(map(len,map(str,row)))<32767
 assert '尾帧' not in row[5] and '站位图' not in row[5]
report={'result':'passed','scope':'26—30集文本、导出读回与格式核对','previous_rows_preserved':24,'appended_rows':len(rows),'data_rows_total':45,'worksheet_count':1,'frozen_panes':b.freeze_panes,'exact_values':True,'fonts_alignment_borders_number_formats_preserved':True,'widths_and_heights_preserved':True,'scene_names_exact_existing_assets':True,'previews_reviewed':['start','middle','end'],'video_validation':'未实测','sha256':hashlib.sha256((BASE/'checked-workbook.xlsx').read_bytes()).hexdigest()}
(BASE/'WORKBOOK_VERIFICATION.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
