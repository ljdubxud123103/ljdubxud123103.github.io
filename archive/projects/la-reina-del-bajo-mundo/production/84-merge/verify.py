from pathlib import Path
from copy import copy
import openpyxl,json,re,hashlib,shutil
base=Path(__file__).resolve().parent
sheet=openpyxl.load_workbook(base/'checked-workbook.xlsx').active
original=Path('/Users/apple/Desktop/前夫黑道酵母_分镜_81-93.xlsx')
ref=openpyxl.load_workbook(original).active
rows=[list(r) for r in sheet.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
expected=json.loads((base/'rows.json').read_text())
norm=lambda r:[None if v=='' else v for v in r]
assert list(map(norm,rows))==list(map(norm,expected))
assert len(rows)==1
src=json.loads((base/'SOURCE_ROWS.json').read_text())
assert rows[0][7]==src[0][7]+'\n'+src[1][7]
assert rows[0][6]==src[1][6]
assert rows[0][4]=='iPhone手机'
prompt=rows[0][5]
assert not any(t in prompt for t in ['米色风衣','@米色风衣','衣凳','披好'])
times=list(map(float,re.findall(r'【时长】([0-9.]+)s',prompt)))
assert times==[3,4,4,4,4,4,4,3] and sum(times)==30
assert prompt.startswith(src[0][5].split('\n\n[场景基调]')[0])
speech=[line.split('：',1)[1][1:-1] for line in src[1][6].splitlines()]
assert len(speech)==4 and all(prompt.count('「'+t+'」')==1 for t in speech)
assert sheet.freeze_panes==ref.freeze_panes
assert [c.value for c in sheet[1]]==[c.value for c in ref[1]]
for ci in range(1,10):
 for prop in ['font','alignment','border','number_format']:assert copy(getattr(sheet.cell(2,ci),prop))==copy(getattr(ref.cell(2,ci),prop))
for col in 'ABCDEFGHI':assert sheet.column_dimensions[col].width==ref.column_dimensions[col].width
assert hashlib.sha256(original.read_bytes()).hexdigest()==json.loads((base.parent/'production-81-continuity/SEPARATE_DELIVERY_CHECK.json').read_text())['original_sha256']
out=Path('/Users/apple/Desktop/84-2A_合并修改稿_30s.xlsx')
shutil.copy2(base/'checked-workbook.xlsx',out)
report={'result':'passed','storyboards':1,'input':['84-2A','84-3A'],'output':'84-2A','seconds':30,'shots':8,'speech_count':4,'VO_count':3,'dialogue_count':1,'source_and_punctuation_preserved':True,'clothing_prop_and_costume_actions_removed':True,'format_preserved':True,'original_workbook_unmodified':True,'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'video_validation':'未实测'}
(base/'WORKBOOK_CHECK.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False))
